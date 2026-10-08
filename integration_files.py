import csv
import json
import traceback
from io import StringIO, BytesIO
from pathlib import Path
from datetime import datetime

# ЧТЕНИЕ ФАЙЛОВ

def _read_csv(file_bytes: bytes) -> list[dict]:
    text = file_bytes.decode("utf-8-sig", errors="replace")
    first_line = text.split("\n", 1)[0]
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    reader = csv.DictReader(StringIO(text), delimiter=delimiter)
    return [dict(row) for row in reader]


def _read_xlsx(file_bytes: bytes) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else f"col_{i}"
               for i, h in enumerate(rows[0])]
    result = []
    for row in rows[1:]:
        if all(c is None for c in row):
            continue
        result.append({
            headers[i]: ("" if v is None else str(v).strip())
            for i, v in enumerate(row) if i < len(headers)
        })
    return result


def _read_json(file_bytes: bytes) -> list[dict]:
    text = file_bytes.decode("utf-8-sig", errors="replace")
    data = json.loads(text)
    if isinstance(data, dict):
        for v in data.values():
            if isinstance(v, list):
                return v
        return [data]
    return data


def read_uploaded_file(filename: str, file_bytes: bytes) -> list[dict]:
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        return _read_csv(file_bytes)
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(file_bytes)
    if ext == ".xls":
        raise ValueError("старый формат .xls — пересохраните файл в .xlsx или .csv")
    if ext == ".json":
        return _read_json(file_bytes)
    raise ValueError(f"формат {ext} не поддерживается (нужен CSV, XLSX или JSON)")

# АВТООПРЕДЕЛЕНИЕ ТИПА

def detect_source_type(filename: str, rows: list[dict]) -> str | None:
    fname = filename.lower()

    name_hints = {
        "hr":        ["hr_", "hr.", "кадр", "employees", "staff", "сотрудник"],
        "warehouse": ["warehouse", "склад", "siz_", "siz.", "карточк", "запас"],
        "training":  ["training", "обучен", "журнал_обуч", "занят"],
        "nfgo":      ["nfgo", "нфго", "формирован", "приказ"],
    }
    for kind, hints in name_hints.items():
        if any(h in fname for h in hints):
            return kind

    if not rows:
        return None
    cols = {str(c).lower().strip() for c in rows[0].keys()}

    def has(*names):
        return all(n.lower() in cols for n in names)

    if has("fio") and (has("department") or has("employee_id")):
        return "hr"
    if has("siz_id") and (has("expiry_date") or has("party_no")):
        return "warehouse"
    if has("topic") and (has("hours") or has("result")):
        return "training"
    if has("formation") and (has("role") or has("employee_id")):
        return "nfgo"

    if "фио" in cols and ("подразделение" in cols or "должность" in cols):
        return "hr"
    if "наименование" in cols and ("срок годности" in cols or "партия" in cols):
        return "warehouse"
    if "тема" in cols and ("часы" in cols or "результат" in cols):
        return "training"
    if "формирование" in cols and ("роль" in cols or "employee_id" in cols):
        return "nfgo"

    return None

# НОРМАЛИЗАЦИЯ

def normalize_fio(fio: str) -> str:
    parts = str(fio).strip().split()
    if len(parts) < 2:
        return str(fio).strip().capitalize()
    return f"{parts[0].capitalize()} " + ".".join(p[0].upper() for p in parts[1:]) + "."


def normalize_department(dept: str) -> str:
    mapping = {
        "служба охраны":    "Служба охраны",
        "отдел охраны":     "Служба охраны",
        "охрана":           "Служба охраны",
        "инженерный отдел": "Инженерный отдел",
        "связь":            "Служба связи",
    }
    return mapping.get(str(dept).strip().lower(), str(dept).strip())


def parse_date(value):
    if not value or str(value).strip() in ("00.00.0000", "None", ""):
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    return None


def _to_int(value, default=0):
    try:
        return int(float(str(value).strip()))
    except (ValueError, TypeError):
        return default


def normalize_hr(rows):
    out = []
    for r in rows:
        emp_id = str(r.get("employee_id", r.get("id", ""))).strip()
        if not emp_id:
            continue
        out.append({
            "employee_id": emp_id,
            "fio":         normalize_fio(r.get("fio", r.get("ФИО", ""))),
            "department":  normalize_department(
                r.get("department", r.get("Подразделение", ""))),
            "position":    str(r.get("position", r.get("Должность", ""))).strip(),
            "face_size":   _to_int(r.get("face_size", r.get("Размер", 2)), 2),
            "status":      str(r.get("status", r.get("Статус", "Работает"))).strip(),
        })
    return out


def normalize_warehouse(rows):
    out = []
    for r in rows:
        siz_id = str(r.get("siz_id", r.get("id", ""))).strip()
        if not siz_id:
            continue
        mfg = parse_date(r.get("manufacture_date", r.get("Дата изготовления", "")))
        exp = parse_date(r.get("expiry_date", r.get("Срок годности", "")))
        if not exp:
            continue
        out.append({
            "siz_id":           siz_id,
            "name":             str(r.get("name", r.get("Наименование", ""))).strip(),
            "party_no":         str(r.get("party_no", r.get("Партия", ""))).strip(),
            "quantity":         _to_int(r.get("quantity", r.get("Количество", 0))),
            "storage":          str(r.get("storage", r.get("Место хранения", ""))).strip(),
            "manufacture_date": mfg.isoformat() if mfg else None,
            "expiry_date":      exp.isoformat(),
            "status":           str(r.get("status", r.get("Статус", "На хранении"))).strip(),
        })
    return out


def normalize_training(rows):
    out = []
    for r in rows:
        emp_id = str(r.get("employee_id", r.get("id", ""))).strip()
        if not emp_id:
            continue
        d = parse_date(r.get("date", r.get("Дата", "")))
        out.append({
            "employee_id": emp_id,
            "topic":       str(r.get("topic", r.get("Тема", ""))).strip(),
            "hours":       _to_int(r.get("hours", r.get("Часы", 0))),
            "result":      str(r.get("result", r.get("Результат", "Не зачтено"))).strip(),
            "date":        d.isoformat() if d else None,
        })
    return out


def normalize_nfgo(rows):
    out = []
    for r in rows:
        emp_id = str(r.get("employee_id", r.get("id", ""))).strip()
        if not emp_id:
            continue
        out.append({
            "employee_id": emp_id,
            "formation":   str(r.get("formation", r.get("Формирование", ""))).strip(),
            "role":        str(r.get("role", r.get("Роль", ""))).strip(),
        })
    return out

# ГЛАВНАЯ ФУНКЦИЯ ИНТЕГРАЦИИ

def integrate_uploaded_auto(files: list[tuple[str, bytes]], log: list) -> dict:
    """Полный цикл: чтение → определение типа → нормализация → очистка →
       связывание. Все сообщения — понятные, без служебных надписей."""

    log.append("Начинаю обработку файлов.")
    log.append(f"Получено файлов: {len(files)}")
    log.append("")

    # Чтение файлов
    log.append("Чтение файлов")
    raw = {"hr": [], "warehouse": [], "training": [], "nfgo": []}
    undetected = []
    failed = []

    for fname, fbytes in files:
        try:
            rows = read_uploaded_file(fname, fbytes)
        except Exception as e:
            reason = str(e).strip() or "неизвестная ошибка"
            log.append(f"  Не удалось прочитать файл «{fname}»: {reason}")
            failed.append((fname, reason))
            continue

        if not rows:
            log.append(f"  Файл «{fname}» пустой — пропускаю.")
            continue

        kind = detect_source_type(fname, rows)
        if kind is None:
            log.append(f"  Файл «{fname}» пропущен: "
                       f"не удалось понять, что это за данные.")
            log.append(f"     Подсказка: переименуйте файл так, чтобы в имени "
                       f"было «hr», «warehouse», «training» или «nfgo», "
                       f"или проверьте названия колонок.")
            undetected.append(fname)
            continue

        label = {"hr": "Сотрудники",
                 "warehouse": "СИЗ на складе",
                 "training": "Журнал обучения",
                 "nfgo": "Участники формирований"}[kind]
        raw[kind].extend(rows)
        log.append(f"  «{fname}» — {label}: прочитано строк {len(rows)}")

    log.append("")
    log.append(f"Итог чтения:")
    log.append(f"  • Сотрудники:            {len(raw['hr'])}")
    log.append(f"  • СИЗ:                   {len(raw['warehouse'])}")
    log.append(f"  • Обучение:              {len(raw['training'])}")
    log.append(f"  • Участники формирований:{len(raw['nfgo'])}")
    log.append("")

    # Нормализация
    log.append("Приведение данных к единому виду")
    hr = normalize_hr(raw["hr"])
    wh = normalize_warehouse(raw["warehouse"])
    tr = normalize_training(raw["training"])
    nf = normalize_nfgo(raw["nfgo"])
    log.append(f"  Сотрудники:            {len(hr)}")
    log.append(f"  СИЗ:                   {len(wh)}")
    log.append(f"  Обучение:              {len(tr)}")
    log.append(f"  Участники формирований:{len(nf)}")
    log.append("")

    # Очистка
    log.append("Удаление дубликатов и мусора")
    seen, hr_clean = set(), []
    for r in hr:
        if r["employee_id"] in seen:
            continue
        seen.add(r["employee_id"])
        hr_clean.append(r)
    log.append(f"  Сотрудники: удалено дубликатов "
               f"{len(hr) - len(hr_clean)}")

    seen, wh_clean = set(), []
    for r in wh:
        if r["siz_id"] in seen:
            continue
        seen.add(r["siz_id"])
        wh_clean.append(r)
    before = len(wh_clean)
    wh_clean = [r for r in wh_clean if r["quantity"] > 0]
    log.append(f"  СИЗ: удалено дубликатов {len(wh) - before}, "
               f"убрано некорректных количеств {before - len(wh_clean)}")
    log.append("")

    # Связывание
    log.append("Связывание данных")
    employees = {e["employee_id"]: e for e in hr_clean}

    nfgo_members = []
    skipped_nfgo = 0
    for m in nf:
        emp = employees.get(m["employee_id"])
        if not emp:
            skipped_nfgo += 1
            continue
        nfgo_members.append({
            "employee_id": m["employee_id"], "fio": emp["fio"],
            "department": emp["department"], "formation": m["formation"],
            "role": m["role"], "status": emp["status"],
        })
    log.append(f"  Участники формирований: сохранено {len(nfgo_members)}, "
               f"пропущено {skipped_nfgo} (нет такого сотрудника в списке)")

    training = []
    skipped_training = 0
    for t in tr:
        emp = employees.get(t["employee_id"])
        if not emp:
            skipped_training += 1
            continue
        training.append({**t, "fio": emp["fio"]})
    log.append(f"  Обучение: сохранено {len(training)}, "
               f"пропущено {skipped_training} (нет такого сотрудника в списке)")
    log.append("")

    # Что в итоге
    log.append("Готово.")
    log.append(f"  • Сотрудников:            {len(employees)}")
    log.append(f"  • Карточек СИЗ:           {len(wh_clean)}")
    log.append(f"  • Участников формирований:{len(nfgo_members)}")
    log.append(f"  • Записей обучения:       {len(training)}")

    if failed:
        log.append("")
        log.append(f"Не удалось прочитать файлов: {len(failed)}")
        for fname, reason in failed:
            log.append(f"  • {fname}: {reason}")

    if undetected:
        log.append("")
        log.append(f"Пропущено файлов с непонятным типом: {len(undetected)}")
        for f in undetected:
            log.append(f"  • {f}")
        log.append("  Что делать: переименуйте файл или проверьте колонки.")

    return {
        "employees":    list(employees.values()),
        "siz_cards":    wh_clean,
        "nfgo_members": nfgo_members,
        "training":     training,
    }